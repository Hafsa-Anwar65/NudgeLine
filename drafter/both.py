"""One run, two drafts per row: a comment for the lead's post AND a personalized email.

Input: a CSV, Excel (.xlsx) or JSON file where each row has the post content, the business
name, the owner's name and an email. Column names are matched loosely (see FIELD_MAP in
emails.py). Draft only: nothing is posted or sent.
"""

import csv
import hashlib
import json
import random
import sys
import time
from pathlib import Path

from . import comments, emails, quality, review
from .common import done_ids, open_writer, prepare_output

OUTPUT_FIELDS = [
    "id", "priority", "intent_score", "language", "name", "company", "email", "post_url", "post_text",
    "comment", "comment_long", "comment_note", "subject", "body", "personalization", "angle", "qa_flags", "status",
]

EMAIL_EXTRA_RULE = (
    "\n\nExtra rule for this run: if the lead details include their public post, mention what "
    "they asked for or talked about in it, in a natural way, as the reason you are writing. "
    "Do not quote it word for word."
)


def _read_rows(path):
    """Return a list of dicts (raw column names) from csv, xlsx or json."""
    p = Path(path)
    if not p.exists():
        sys.exit(f"Input file not found: {path}")
    suffix = p.suffix.lower()

    if suffix == ".csv":
        with p.open(newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))

    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("To read Excel files run: pip install openpyxl  (or save the sheet as CSV)")
        ws = load_workbook(p, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h) if h is not None else "" for h in rows[0]]
        return [
            {header[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(header)}
            for r in rows[1:]
        ]

    if suffix == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            for key in ("posts", "results", "data", "items", "leads"):
                if isinstance(data.get(key), list):
                    data = data[key]
                    break
            else:
                data = [data]
        return [r for r in data if isinstance(r, dict)]

    sys.exit("Use a .csv, .xlsx or .json file for this command.")


def load_records(path):
    records, seen = [], set()
    for raw in _read_rows(path):
        row = {emails.norm_key(k): ("" if v is None else str(v).strip()) for k, v in raw.items()}
        rec = {}
        for field, aliases in emails.FIELD_MAP.items():
            rec[field] = next((row[a] for a in aliases if row.get(a)), "")
        rec["email"] = rec["email"].lower()
        if not rec["notes"] and not rec["email"]:
            continue  # nothing to work with
        key = rec["email"] + "|" + rec["notes"][:200] + "|" + rec["company"]
        rec["id"] = hashlib.md5(key.encode("utf-8")).hexdigest()[:12]
        if rec["id"] in seen:
            continue
        seen.add(rec["id"])
        records.append(rec)
    return records


def run(args, profile, llm):
    template = emails.load_template(getattr(args, "template", None))
    records = load_records(args.input)
    prepare_output(args.output, OUTPUT_FIELDS)
    done = done_ids(args.output, "id")
    todo = [r for r in records if r["id"] not in done]
    if args.limit:
        todo = todo[: args.limit]
    print(f"{len(records)} rows loaded, {len(done)} already done, {len(todo)} to process.")
    if template:
        print("Using your email template: the AI only fills the placeholders.")

    length = profile.get("comment_length", "short")
    comment_system = comments.build_system_prompt(profile, length)
    email_system = emails.make_system(profile, template)
    if not template:
        email_system += EMAIL_EXTRA_RULE

    f, writer = open_writer(args.output, OUTPUT_FIELDS)
    with f:
        for i, rec in enumerate(todo, 1):
            row = {
                "id": rec["id"], "priority": "", "intent_score": "", "language": "",
                "name": rec["name"], "company": rec["company"],
                "email": rec["email"], "post_url": rec["post_url"],
                "post_text": rec["notes"][:1500].replace("\n", " "),
                "comment": "", "comment_long": "", "comment_note": "", "subject": "", "body": "",
                "personalization": "", "angle": "", "qa_flags": "", "status": "",
            }
            failed = False
            flags = []

            # 1) Comment on the post
            if rec["notes"]:
                style = random.choice(comments.STYLE_HINTS)
                user_msg = (
                    f"Style for this comment: {style}\n\n"
                    f"Post author: {rec['name'] or 'unknown'}\n"
                    f"Post text:\n{rec['notes'][:2000]}"
                )
                fallback = {
                    "relevant": True, "intent_score": 3, "language": "English",
                    "comment": "[dry-run] a drafted comment would appear here",
                    "reason": "dry run, no API call made",
                }
                if length == "long":
                    fallback["comment"] = "[dry-run] a longer drafted comment would appear here, with more detail about how you would help."
                if length == "both":
                    fallback["comment_long"] = "[dry-run] a longer drafted comment would appear here, with more detail about how you would help."
                info = comments.interpret(llm.ask_json(comment_system, user_msg, fallback, max_tokens=comments.token_budget(length)), length)
                if "error" in info:
                    failed = True
                    row["comment_note"] = f"error: {info['error']}"
                else:
                    row["priority"], row["language"] = info["priority"], info["language"]
                    row["intent_score"] = info["intent"] or ""
                    flags += info["flags"]
                    if info["relevant"] and info["comment"]:
                        row["comment"], row["comment_note"] = info["comment"], info["reason"]
                        row["comment_long"] = info.get("comment_long", "")
                    elif info["relevant"]:
                        failed = True
                        row["comment_note"] = "error: empty comment"
                    else:
                        row["comment_note"] = "skipped: " + (info["reason"] or "not relevant")
            else:
                row["comment_note"] = "no post text"

            # 2) Email to the owner
            if rec["email"] and emails.EMAIL_RE.match(rec["email"]):
                site = ""
                if args.fetch_sites and rec["website"] and not llm.dry_run:
                    site = emails.fetch_site_summary(rec["website"])
                fallback = {
                    "subject": "[dry-run] subject line",
                    "body": "[dry-run] a personalized email body would appear here.",
                    "personalization": "low",
                    "angle": "dry run, no API call made",
                }
                message = emails.build_user_message(rec, site, notes_label="Their public post")
                res = llm.ask_json(email_system, message, fallback, max_tokens=700)
                if "_error" in res or not res.get("body"):
                    failed = True
                    row["angle"] = f"error: {res.get('_error', 'empty response')}"
                else:
                    source = " ".join([rec["name"], rec["company"], rec["notes"], site])
                    subject, full_body, email_flags = emails.finish_email(res, profile, source, template)
                    row["subject"], row["body"] = subject, full_body
                    row["personalization"] = res.get("personalization", "")
                    row["angle"] = res.get("angle", "")
                    flags += email_flags
            else:
                row["angle"] = "no valid email, email draft skipped"

            row["qa_flags"] = "; ".join(flags)
            if failed:
                row["status"] = "error"
            elif flags:
                row["status"] = "needs_edit"
            else:
                row["status"] = "draft" if (row["comment"] or row["body"]) else "skipped"
            writer.writerow(row)
            f.flush()
            label = rec["company"] or rec["name"] or rec["email"] or rec["id"]
            print(f"[{i}/{len(todo)}] {row['status']}: {label}")
            if not llm.dry_run:
                time.sleep(args.gap)

    review.sort_by_priority(args.output)
    page = review.write_html(args.output)
    print(f"Done. Open {page} to review the drafts.")
    print("Edit the CSV if you like (set status to 'draft' for rows you approve), then run:")
    print(f'  python main.py approve --input "{args.output}"')
